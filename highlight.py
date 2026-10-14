from openpyxl import load_workbook, Workbook
from copy import copy

# Open the first Excel file
wb = load_workbook(r"C:\Users\Kavitha\OneDrive\Documents\emp7.xlsx")
ws = wb.active

# Create the second Excel file
new_wb = Workbook()
new_ws = new_wb.active

# Copy column names (first row)
for cell in ws[1]:
    new_cell = new_ws.cell(row=1, column=cell.column, value=cell.value)
    new_cell.font = copy(cell.font)
    new_cell.fill = copy(cell.fill)

# Copy highlighted employee rows
output_row = 2

for row in ws.iter_rows(min_row=2):
    if any(cell.fill.fill_type == "solid" for cell in row):
        for cell in row:
            new_cell = new_ws.cell(
                row=output_row,
                column=cell.column,
                value=cell.value
            )
            new_cell.fill = copy(cell.fill)
            new_cell.font = copy(cell.font)
            new_cell.border = copy(cell.border)
            new_cell.alignment = copy(cell.alignment)

        output_row += 1

# Save the second Excel file
new_wb.save(r"C:\Users\Kavitha\OneDrive\Documents\emp8.xlsx")

print("Column names and highlighted rows copied successfully!")