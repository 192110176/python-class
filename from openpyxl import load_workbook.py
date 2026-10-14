from openpyxl import load_workbook

# Step 1: Open the source Excel file
source_file = "Source.xlsx"
source_wb = load_workbook(source_file, data_only=True)
source_ws = source_wb["Sheet1"]

# Step 2: Open the destination Excel file
destination_file = "Destination.xlsx"
destination_wb = load_workbook(destination_file)
destination_ws = destination_wb["Sheet1"]

# Step 3: Read data from source and copy it
for row in source_ws.iter_rows():
    for cell in row:
        destination_ws[cell.coordinate] = cell.value

# Step 4: Save the destination Excel file
destination_wb.save(destination_file)

# Step 5: Close workbooks
source_wb.close()
destination_wb.close()

print("Data successfully copied to Destination.xlsx")